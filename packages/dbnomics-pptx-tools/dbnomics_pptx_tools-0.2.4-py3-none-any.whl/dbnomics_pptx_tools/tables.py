from dataclasses import dataclass
from types import ModuleType
from typing import Iterable, Iterator, TypeAlias, cast

import daiquiri
from parsy import ParseError, any_char, char_from, regex, seq, string
from pptx.table import Table, _Cell, _Row, _RowCollection
from tabulate import tabulate  # type: ignore

from dbnomics_pptx_tools.formatters import format_number
from dbnomics_pptx_tools.metadata import TableSpec
from dbnomics_pptx_tools.repo import SeriesRepo

__all__ = ["extract_table_zones", "format_table", "replace_cell_text", "update_table"]

logger = daiquiri.getLogger(__name__)


@dataclass
class AnnualPeriod:
    year: int

    def __str__(self) -> str:
        return str(self.year)


@dataclass
class QuarterlyPeriod:
    year: int
    quarter: int

    def __str__(self) -> str:
        return f"{self.year}-Q{self.quarter}"


Period: TypeAlias = AnnualPeriod | QuarterlyPeriod


@dataclass
class TableZones:
    first_data_row_index: int
    periods: list[Period]


def extract_table_zones(table: Table, *, adhoc_module: ModuleType | None, table_name: str) -> TableZones | None:
    header_row_index = find_header_row_index(table)
    if header_row_index is None:
        logger.debug("Could not find the row corresponding to the table header")
        return None
    periods = parse_header_periods(
        table, adhoc_module=adhoc_module, header_row_index=header_row_index, table_name=table_name
    )
    if periods is None:
        logger.debug("Could not parse the periods in the table header")
        return None
    first_data_row_index = find_first_data_row_index(table, header_row_index=header_row_index)
    if first_data_row_index is None:
        logger.debug("Could not find the first data row of the table")
        return None
    return TableZones(first_data_row_index=first_data_row_index, periods=periods)


def find_first_data_row_index(table: Table, *, header_row_index: int) -> int | None:
    rows = list(cast(Iterable[_Row], table.rows))
    start = header_row_index + 1
    for row_index, row in enumerate(rows[start:], start=start):
        first_cell_text = row.cells[0].text.strip()
        if not first_cell_text:
            continue
        return row_index
    return None


def find_header_row_index(table: Table) -> int | None:
    accepted_header_first_cells = {"Country"}
    rows = cast(Iterable[_Row], table.rows)
    for row_index, row in enumerate(rows):
        if row.cells[0].is_merge_origin:
            if row_index == 0:
                continue
            raise NotImplementedError()
        if row.cells[0].text in accepted_header_first_cells:
            return row_index
    return None


def format_table(table: Table) -> str:
    rows = cast(Iterable[_Row], table.rows)
    tabular_data = [[cell.text for cell in row.cells] for row in rows]
    return tabulate(tabular_data)


def get_period_cells(header_cells: list[_Cell], *, adhoc_module: ModuleType | None, table_name: str) -> list[_Cell]:
    period_cells = header_cells[1:]
    if adhoc_module is not None and hasattr(adhoc_module, "filter_period_cells"):
        filtered_period_cells = adhoc_module.filter_period_cells(period_cells, table_name=table_name)
        if filtered_period_cells is not None:
            return filtered_period_cells
    return period_cells


def iter_table_data_rows(table: Table, *, table_zones: TableZones) -> Iterator[tuple[int, str]]:
    rows = list(cast(Iterable[_Row], table.rows))
    first_data_row_index = table_zones.first_data_row_index
    for row_index, row in enumerate(rows[first_data_row_index:], start=first_data_row_index):
        row_label = row.cells[0].text.strip()
        yield row_index, row_label


def parse_header_period(text: str) -> Period:
    def to_full_year(x: str) -> int:
        return int(x) + 2000

    full_year = regex(r"[0-9]{4}").desc("4 digit year")
    short_year = regex(r"[0-9]{2}").desc("2 digit year")
    annual_period = full_year.map(int)
    quarterly_period = seq(year=short_year.map(to_full_year), quarter=(string("Q") >> char_from("1234").map(int)))
    period = annual_period.map(AnnualPeriod) | quarterly_period.combine_dict(QuarterlyPeriod)
    attribute = string("(") >> any_char << string(")")
    period_header = period << attribute.optional()
    return period_header.parse(text)


def parse_header_periods(
    table: Table, *, adhoc_module: ModuleType | None, header_row_index: int, table_name: str
) -> list[Period] | None:
    rows = cast(_RowCollection, table.rows)
    header_cells = list(rows[header_row_index].cells)
    period_cells = get_period_cells(header_cells, adhoc_module=adhoc_module, table_name=table_name)
    period_cells_text = [cell.text for cell in period_cells]
    logger.debug("Parsing period cells: %r", period_cells_text)

    has_errors = False

    def iter_parsed():
        nonlocal has_errors
        for text_pos, text in enumerate(period_cells_text):
            try:
                yield parse_header_period(text)
            except ParseError:
                logger.exception("Could not parse period at index %d from text %r", text_pos, text)
                has_errors = True

    return None if has_errors else list(iter_parsed())


def replace_cell_text(cell: _Cell, text: str):
    runs = cell.text_frame.paragraphs[0].runs
    if not runs:
        raise ValueError("No text in cell, do not know which format (font, alignment, etc.) to use")
    runs[0].text = text


def update_table(
    table: Table,
    *,
    adhoc_module: ModuleType | None,
    repo: SeriesRepo,
    table_name: str,
    table_spec: TableSpec,
    table_zones: TableZones,
):
    for row_index, row_label in iter_table_data_rows(table, table_zones=table_zones):
        for col_index, period in enumerate(table_zones.periods, start=1):
            update_table_cell(
                table,
                col_index=col_index,
                period=str(period),
                repo=repo,
                row_index=row_index,
                row_label=row_label,
                table_spec=table_spec,
            )
    if adhoc_module is not None and hasattr(adhoc_module, "process_adhoc_table"):
        adhoc_module.process_adhoc_table(
            table,
            repo=repo,
            table_name=table_name,
            table_spec=table_spec,
            table_zones=table_zones,
        )


def update_table_cell(
    table: Table,
    *,
    col_index: int,
    period: str,
    repo: SeriesRepo,
    row_index: int,
    row_label: str,
    table_spec: TableSpec,
):
    series_id = table_spec.find_series_id_by_name(row_label)
    if series_id is None:
        logger.warning("Could not find the series ID from the row label %r, ignoring row", row_label)
        return
    df = repo.load(series_id)
    observations = df[df["original_period"] == period].value
    cell = table.cell(row_index, col_index)
    if observations.empty:
        dash = "–"
        logger.debug(
            "Period %r requested for table, but not found in series %r, fallback to %r in table cell",
            period,
            series_id,
            dash,
        )
        replace_cell_text(cell, dash)
    elif len(observations) > 1:
        logger.warning("Many observations found for period %r in series %r, ignoring period", period, series_id)
    else:
        observation = cast(float, observations.values[0])
        replace_cell_text(cell, format_number(observation))
